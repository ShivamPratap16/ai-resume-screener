"""Generates a synthetic, varied resume set (plus edge-case files) for exercising the pipeline.

All names, emails and GitHub handles are fictional.
"""
import json
import random
import shutil
from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "fixtures" / "sample_resumes"
MANIFEST = ROOT / "tests" / "fixtures" / "sample_manifest.json"

P = {
    "deep_agent": ("Multi-Agent Support Copilot | Python, LangGraph, FastAPI, PostgreSQL, Redis", [
        "Built a LangGraph multi-agent workflow (planner, retriever, executor) with tool calling against order and refund APIs",
        "Implemented hybrid search with pgvector embeddings and re-ranking over 40k support tickets",
        "Persisted agent state with PostgreSQL checkpoints; Redis caching cut p95 latency from 4.1s to 1.3s",
        "Added an evaluation pipeline with RAGAS and a 200-question golden set that runs in GitHub Actions",
        "Async FastAPI service with retries, timeouts and structured logging; deployed on GCP Cloud Run using Docker",
    ]),
    "rag_solid": ("Course Notes RAG Assistant — Python, LlamaIndex, Qdrant, FastAPI", [
        "Designed an ingestion pipeline that parses lecture PDFs, applies semantic chunking and stores embeddings in Qdrant",
        "Exposed retrieval and answer endpoints through FastAPI with async request handling",
        "Evaluated answer faithfulness with an LLM-as-judge harness over 150 labelled questions",
    ]),
    "adk_agent": ("Research Agent with Google ADK | Python, Gemini, Firestore", [
        "Built a research agent with Google ADK that plans sub-queries and uses tool calling for web search and citation lookup",
        "Kept conversation memory and session state in Firestore so long research tasks can resume",
        "Wrote pytest unit tests for each tool and deployed the agent on GCP Cloud Run",
    ]),
    "sql_agent": ("Text-to-SQL Agent | Python, LangChain, PostgreSQL", [
        "Built an SQL agent with tool calling that answers analytics questions over a PostgreSQL warehouse",
        "Added guardrails that validate generated SQL against the schema and block write queries",
        "Logged every agent step for observability and replayed failures in an evaluation set",
    ]),
    "mcp": ("MCP Tool Server for Code Search | Python, FastAPI", [
        "Implemented a Model Context Protocol server that exposes repository search and file tools to LLM agents",
        "Added rate limiting, caching of repository indexes and integration tests with pytest",
    ]),
    "eval": ("LLM Evaluation Harness | Python, DeepEval, asyncio", [
        "Built an evaluation pipeline comparing prompts and models on 300 test cases using DeepEval metrics",
        "Ran model calls concurrently with asyncio and a semaphore; results stored in SQLite for regression tracking",
    ]),
    "fullstack_rag": ("DocuMind — document Q&A platform | Next.js, FastAPI, LangChain, Pinecone, Docker", [
        "Built an end-to-end RAG product: Next.js frontend, FastAPI backend and Pinecone vector search",
        "Implemented document ingestion with chunking, embeddings and metadata filters per workspace",
        "Streamed answers over websockets and stored chat history in PostgreSQL; deployed with Docker on AWS",
    ]),
    "thin": ("AI Chatbot | Python, Streamlit, OpenAI API", [
        "Simple chatbot that sends the user prompt to the OpenAI API and displays the response in Streamlit",
    ]),
    "thin2": ("Resume Rewriter using GPT-4 | Python, Flask", [
        "Flask app that calls the OpenAI API to rewrite resume bullet points in a more professional tone",
    ]),
    "tutorial": ("Chat with PDF (LangChain)", [
        "Followed a YouTube tutorial to build a chat with PDF app using LangChain",
    ]),
    "classic_ml": ("House Price Prediction | Python, scikit-learn, pandas", [
        "Trained gradient boosting and linear regression models on 20k listings with feature engineering in pandas",
        "Tuned hyperparameters with cross-validation and reached an R2 score of 0.89",
    ]),
    "cnn": ("Plant Disease Classifier | Python, PyTorch", [
        "Trained a ResNet CNN on 50k leaf images with data augmentation, reaching 96% accuracy",
    ]),
    "web_py": ("Inventory Management API | Python, Django, PostgreSQL", [
        "Developed REST APIs with Django REST Framework for stock tracking across 12 warehouses",
        "Wrote unit tests with pytest and added Celery jobs for nightly stock reconciliation",
    ]),
    "data_eng": ("Sales ETL Pipeline | Python, Airflow, BigQuery", [
        "Built Airflow DAGs that ingest sales CSVs, clean them with pandas and load them into BigQuery",
        "Added data quality checks and monitoring alerts for failed loads",
    ]),
    "java": ("Library Management System | Java, Spring Boot, MySQL", [
        "Developed Spring Boot services for issuing and returning books with MySQL persistence",
        "Implemented JWT authentication and role-based access control",
    ]),
    "mern": ("E-commerce Platform | React, Node.js, Express, MongoDB", [
        "Built a full-stack shop with React, Redux and an Express REST API backed by MongoDB",
        "Integrated Stripe payments and deployed the frontend on Vercel",
    ]),
    "js_ai": ("AI Code Reviewer | Next.js, TypeScript, LangChain.js, OpenAI", [
        "Built a Next.js app using LangChain.js agents with tool calling to review pull requests",
        "Implemented embeddings-based search over repository files with a vector store in Supabase",
    ]),
}

E = {
    "ai_intern": ("AI Engineering Intern, Finlytics | May 2025 – Aug 2025", [
        "Built a RAG pipeline in Python over 8k policy documents using LangChain and FAISS",
        "Added Redis caching and retries with exponential backoff for LLM calls, cutting cost by 30%",
        "Wrote pytest unit and integration tests for the retrieval layer",
    ]),
    "backend_intern": ("Backend Intern, ShopStack | Jan 2025 – Jun 2025", [
        "Developed async FastAPI services with PostgreSQL and SQLAlchemy handling 2M requests/day",
        "Containerized services with Docker and set up CI/CD with GitHub Actions",
    ]),
    "java_intern": ("Software Engineering Intern, BankCore | Jun 2024 – Aug 2024", [
        "Developed Spring Boot microservices in Java for account statements",
    ]),
    "frontend_intern": ("Frontend Intern, PixelWorks | Jun 2024 – Dec 2024", [
        "Built reusable React and TypeScript components for a design system used by 5 teams",
    ]),
    "ds_intern": ("Data Science Intern, RetailCo | Jan 2024 – Jun 2024", [
        "Built a churn prediction model with scikit-learn and pandas, presented insights to leadership",
    ]),
}

SKILLS = {
    "ai_py": "Python, FastAPI, LangGraph, LangChain, PostgreSQL, Redis, Docker, GCP, Git",
    "rag_py": "Python, LlamaIndex, LangChain, Qdrant, FastAPI, SQL, Docker, AWS",
    "py_basic": "Python, Flask, Streamlit, SQL, Git, HTML, CSS",
    "py_ml": "Python, scikit-learn, pandas, NumPy, PyTorch, Matplotlib, SQL",
    "py_ml_llm": "Python, pandas, LangChain, OpenAI API, scikit-learn, SQL",
    "py_backend": "Python, Django, FastAPI, PostgreSQL, Redis, Celery, Docker, AWS",
    "py_backend_llm": "Python, Django, PostgreSQL, Docker, LangChain, LlamaIndex, Pinecone",
    "java": "Java, Spring Boot, MySQL, Hibernate, Git, Jenkins",
    "mern": "JavaScript, TypeScript, React, Next.js, Node.js, Express, MongoDB",
    "js_ai": "TypeScript, Next.js, React, LangChain.js, OpenAI API, Supabase",
    "data": "Python, SQL, Airflow, BigQuery, pandas, dbt, GCP",
}

HEADINGS = [
    {"summary": "SUMMARY", "skills": "TECHNICAL SKILLS", "experience": "EXPERIENCE", "projects": "PROJECTS", "education": "EDUCATION"},
    {"summary": "About Me", "skills": "Tech Stack", "experience": "Internships", "projects": "Things I've Built", "education": "Academics"},
    {"summary": "Profile", "skills": "Core Competencies", "experience": "Professional Experience", "projects": "Academic Projects", "education": "Education"},
]

# archetype -> (count, expected eligible, skills key, project keys, experience keys, summary)
ARCHETYPES = {
    "strong_agentic": (4, True, "ai_py", [["deep_agent", "rag_solid"], ["deep_agent", "adk_agent"], ["deep_agent", "sql_agent"], ["deep_agent", "mcp"]], ["ai_intern", "backend_intern"], "Engineer building agentic AI systems in Python."),
    "solid_rag": (4, True, "rag_py", [["rag_solid", "web_py"], ["fullstack_rag", "web_py"], ["rag_solid", "eval"], ["fullstack_rag", "classic_ml"]], ["backend_intern", None], "Backend-focused student working on retrieval systems."),
    "agent_mid": (4, True, "py_ml_llm", [["adk_agent", "classic_ml"], ["sql_agent", "cnn"], ["mcp", "classic_ml"], ["eval", "web_py"]], [None, "ds_intern"], "CS undergraduate interested in AI agents."),
    "thin_wrapper": (4, True, "py_basic", [["thin", "web_py"], ["thin2", "classic_ml"], ["thin", "thin2"], ["thin2", "web_py"]], [None], "Enthusiastic developer exploring GenAI."),
    "tutorial": (3, True, "py_ml_llm", [["tutorial", "classic_ml"], ["tutorial", "cnn"], ["tutorial", "web_py"]], [None, "ds_intern"], "Aspiring ML engineer."),
    "skills_only_ai": (3, True, "py_backend_llm", [["web_py", "data_eng"], ["web_py"], ["data_eng", "classic_ml"]], ["backend_intern"], "Python developer."),
    "python_no_ai": (3, False, "py_backend", [["web_py", "data_eng"], ["web_py"], ["data_eng"]], ["backend_intern"], "Backend developer who enjoys clean APIs."),
    "java_only": (3, False, "java", [["java"], ["java", "mern"], ["java"]], ["java_intern"], "Java developer focused on enterprise backends."),
    "mern_only": (3, False, "mern", [["mern"], ["mern", "java"], ["mern"]], ["frontend_intern"], "Full-stack JavaScript developer."),
    "js_ai_no_python": (3, False, "js_ai", [["js_ai", "mern"], ["js_ai"], ["js_ai", "mern"]], ["frontend_intern"], "TypeScript engineer building AI products."),
    "classic_ml_only": (3, False, "py_ml", [["classic_ml", "cnn"], ["cnn"], ["classic_ml"]], ["ds_intern"], "Data science student."),
    "learning_python": (2, False, "js_ai", [["js_ai"], ["js_ai", "mern"]], ["frontend_intern"], "TypeScript developer. Currently learning Python for backend work."),
}

FIRST = ["Asha", "Rohan", "Priya", "Arjun", "Meera", "Kabir", "Ishita", "Vikram", "Neha", "Aditya", "Sana", "Karthik",
         "Divya", "Rahul", "Ananya", "Farhan", "Pooja", "Siddharth", "Tanvi", "Nikhil", "Riya", "Varun", "Aisha", "Dev",
         "Lakshmi", "Manav", "Zoya", "Harsh", "Kavya", "Omkar", "Simran", "Yash", "Nandini", "Aman", "Bhavna", "Kunal",
         "Esha", "Gaurav", "Hina", "Jay", "Mitali", "Pranav", "Shreya", "Tarun", "Uma", "Vivek", "Aarav", "Diya"]
LAST = ["Rao", "Mehta", "Iyer", "Sharma", "Nair", "Singh", "Gupta", "Reddy", "Kulkarni", "Das", "Khan", "Menon", "Joshi",
        "Pillai", "Chopra", "Bose", "Verma", "Patel", "Saxena", "Shetty"]


def wrap(text: str, width: float, size: int) -> list[str]:
    return simpleSplit(text, "Helvetica", size, width)


def write_pdf(path: Path, lines: list[tuple[str, str]]) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    _, height = A4
    y, left, width = height - 50, 50, 495
    for kind, text in lines:
        font, size, indent = {"name": ("Helvetica-Bold", 16, 0), "h": ("Helvetica-Bold", 11, 0),
                              "title": ("Helvetica-Bold", 10, 0), "bullet": ("Helvetica", 9.5, 12),
                              "text": ("Helvetica", 9.5, 0)}[kind]
        if kind == "h":
            y -= 6
        prefix = "• " if kind == "bullet" else ""
        for i, chunk in enumerate(wrap(prefix + text, width - indent, size)):
            if y < 50:
                c.showPage()
                y = height - 50
            c.setFont(font, size)
            c.drawString(left + indent + (8 if i and kind == "bullet" else 0), y, chunk)
            y -= size + 4
    c.save()


def build_lines(name, email, github, summary, skills, projects, experience, headings, edu):
    contact = f"{email} | +91 98{random.randint(10000000, 99999999)}" + (f" | github.com/{github}" if github else "")
    lines = [("name", name), ("text", contact), ("h", headings["summary"]), ("text", summary),
             ("h", headings["skills"]), ("text", skills)]
    order = ["experience", "projects"] if random.random() < 0.5 else ["projects", "experience"]
    for section in order:
        entries = experience if section == "experience" else projects
        if not entries:
            continue
        lines.append(("h", headings[section]))
        for title, bullets in entries:
            lines.append(("title", title))
            lines += [("bullet", b) for b in bullets]
    lines += [("h", headings["education"]), ("text", edu)]
    return lines


def main() -> None:
    random.seed(7)
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    names = random.sample([f"{f} {l}" for f in FIRST for l in LAST], 60)
    manifest = {}
    n = 0

    for archetype, (count, eligible, skills_key, project_sets, exp_options, summary) in ARCHETYPES.items():
        for i in range(count):
            n += 1
            name = names.pop()
            handle = f"zz-sample-{name.split()[0].lower()}-{n:02d}" if random.random() < 0.6 else None
            email = f"{name.lower().replace(' ', '.')}{n}@example.com"
            exp_key = exp_options[i % len(exp_options)]
            lines = build_lines(
                name, email, handle, summary, SKILLS[skills_key],
                [P[k] for k in project_sets[i % len(project_sets)]],
                [E[exp_key]] if exp_key else [],
                HEADINGS[n % len(HEADINGS)],
                f"B.Tech Computer Science, Example Institute of Technology | 2022 – 2026 | CGPA {random.uniform(7, 9.6):.1f}",
            )
            filename = f"candidate_{n:02d}.pdf"
            write_pdf(OUT / filename, lines)
            manifest[filename] = {"archetype": archetype, "eligible": eligible, "name": name}

    narrative = [
        ("narrative_strong", True, "Kiran Sood", [
            "I am a final-year student who builds agentic AI systems in Python.",
            "Last summer I built a LangGraph agent that triages customer emails using tool calling against a CRM, with conversation memory stored in PostgreSQL.",
            "I also implemented a RAG pipeline with embeddings in pgvector and wrote an evaluation suite with RAGAS, served through an async FastAPI backend deployed with Docker on GCP.",
        ]),
        ("narrative_strong", True, "Leela Thomas", [
            "Python developer who shipped a retrieval-augmented support bot for a college helpdesk.",
            "I designed the ingestion pipeline with chunking and embeddings in ChromaDB, and built FastAPI endpoints with Redis caching.",
            "I wrote pytest tests and added retries and timeouts around LLM calls.",
        ]),
        ("narrative_weak", False, "Mohit Arora", [
            "Motivated web developer with experience in React, Node.js and MongoDB.",
            "I built a portfolio website and a to-do app and deployed them on Vercel.",
        ]),
    ]
    for archetype, eligible, name, paragraphs in narrative:
        n += 1
        filename = f"candidate_{n:02d}.pdf"
        email = f"{name.lower().replace(' ', '.')}@example.com"
        write_pdf(OUT / filename, [("name", name), ("text", email)] + [("text", p) for p in paragraphs])
        manifest[filename] = {"archetype": archetype, "eligible": eligible, "name": name}

    for i, (key, eligible) in enumerate([("fullstack_rag", True), ("deep_agent", True)]):
        n += 1
        name = names.pop()
        filename = f"candidate_{n:02d}.pdf"
        lines = build_lines(name, f"{name.lower().replace(' ', '.')}@example.com", f"zz-sample-fs-{n:02d}",
                            "Full-stack engineer shipping AI products end to end.",
                            "Python, FastAPI, Next.js, React, TypeScript, LangChain, Pinecone, PostgreSQL, Docker, GCP",
                            [P[key], P["mern"]], [E["backend_intern"]], HEADINGS[0], "B.Tech IT, Example University | 2026")
        write_pdf(OUT / filename, lines)
        manifest[filename] = {"archetype": "fullstack_ai", "eligible": eligible, "name": name}

    # DOCX and TXT resumes (bonus formats).
    document = docx.Document()
    document.add_heading("Nisha Varghese", 0)
    document.add_paragraph("nisha.varghese@example.com | github.com/zz-sample-nisha-docx")
    document.add_heading("Skills", 1)
    document.add_paragraph("Python, LangGraph, FastAPI, Redis, PostgreSQL, Docker")
    document.add_heading("Projects", 1)
    document.add_paragraph("Travel Planner Agent | Python, LangGraph, FastAPI")
    for bullet in [
        "Built a planning agent with tool calling for flight and hotel search APIs and stateful checkpoints",
        "Added Redis caching of tool results and an evaluation set of 80 itineraries",
    ]:
        document.add_paragraph(bullet, style="List Bullet")
    document.save(OUT / "nisha_varghese.docx")
    manifest["nisha_varghese.docx"] = {"archetype": "docx_agentic", "eligible": True, "name": "Nisha Varghese"}

    (OUT / "farid_ahmed.txt").write_text(
        "Farid Ahmed\nfarid.ahmed@example.com\n\nSKILLS: Python, Flask, OpenAI API, SQL\n\nPROJECTS\n"
        "Email Summarizer | Python, Flask\n- Flask wrapper that calls the OpenAI API to summarize emails pasted by the user\n"
        "\nEDUCATION\nBSc Computer Science, 2026\n",
        encoding="utf-8",
    )
    manifest["farid_ahmed.txt"] = {"archetype": "txt_thin_wrapper", "eligible": True, "name": "Farid Ahmed"}

    # Edge cases that must not crash the batch.
    (OUT / "corrupted_resume.pdf").write_bytes(b"%PDF-1.4\nthis is not really a pdf\x00\x01\x02")
    manifest["corrupted_resume.pdf"] = {"archetype": "corrupt", "status": "failed"}

    c = canvas.Canvas(str(OUT / "scanned_resume.pdf"), pagesize=A4)
    c.rect(100, 400, 300, 200, fill=1)
    c.save()
    manifest["scanned_resume.pdf"] = {"archetype": "image_only", "status": "failed"}

    shutil.copy(OUT / "candidate_01.pdf", OUT / "candidate_01_copy.pdf")
    manifest["candidate_01_copy.pdf"] = {"archetype": "duplicate_file", "status": "duplicate"}

    first = manifest["candidate_02.pdf"]["name"]
    write_pdf(OUT / "candidate_02_resubmission.pdf", [
        ("name", first), ("text", f"{first.lower().replace(' ', '.')}2@example.com"),
        ("h", "SKILLS"), ("text", "Python, LangChain, FastAPI"),
        ("h", "PROJECTS"), ("title", "Updated resume version"), ("bullet", "Same candidate submitting a second time with a tweaked layout and wording."),
    ])
    manifest["candidate_02_resubmission.pdf"] = {"archetype": "duplicate_email", "status": "duplicate"}

    (OUT / "profile_photo.png").write_bytes(b"\x89PNG\r\n\x1a\nnot-a-resume")
    manifest["profile_photo.png"] = {"archetype": "unsupported", "status": "failed"}

    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"Wrote {len(manifest)} files to {OUT}")


if __name__ == "__main__":
    main()
