import hashlib
from pathlib import Path


class ExtractionError(Exception):
    pass


def discover(input_dir: Path) -> list[Path]:
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")
    return sorted(p for p in input_dir.iterdir() if p.is_file() and not p.name.startswith("."))


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_text(path: Path, data: bytes) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _pdf_text(data)
        if suffix == ".docx":
            return _docx_text(data)
        if suffix == ".txt":
            return data.decode("utf-8", errors="replace")
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not read {suffix} file: {exc}") from exc
    raise ExtractionError(f"Unsupported file type: {suffix or 'none'}")


def _pdf_text(data: bytes) -> str:
    import pymupdf

    pymupdf.TOOLS.mupdf_display_errors(False)
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ExtractionError("PDF is password protected")
        # Content-stream order keeps two-column resumes intact; sort=True interleaves the columns.
        text = "\n".join(page.get_text("text") for page in doc)
        # "GitHub"/"LinkedIn" are often clickable labels whose URL only exists as a link annotation.
        links = [link["uri"] for page in doc for link in page.get_links() if link.get("uri")]
    return text + ("\n" + "\n".join(links) if links else "")


def _docx_text(data: bytes) -> str:
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)
